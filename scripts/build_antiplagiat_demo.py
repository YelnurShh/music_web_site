from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path('/Users/elnrsahar/Desktop/music_website/antiplagiat_anyktama_ulgisi.docx')


def set_font(run, size=12, bold=False, color='000000'):
    run.font.name = 'Times New Roman'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn('w:shd'))
    if node is None:
        node = OxmlElement('w:shd')
        props.append(node)
    node.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=150, start=160, bottom=150, end=160):
    props = cell._tc.get_or_add_tcPr()
    tc_mar = props.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        props.append(tc_mar)
    for name, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{name}'))
        if node is None:
            node = OxmlElement(f'w:{name}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.59)
sec.page_height = Cm(27.94)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)
sec.top_margin = Cm(1.8)
sec.bottom_margin = Cm(1.8)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(6)

title = doc.styles['Title']
title.font.name = 'Times New Roman'
title._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
title.font.size = Pt(16)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(6)
title_ppr = title.element.get_or_add_pPr()
border = title_ppr.find(qn('w:pBdr'))
if border is not None:
    title_ppr.remove(border)

# Unmissable status label. This prevents the demonstration from being mistaken for evidence.
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run('ҮЛГІ — РЕСМИ ЕМЕС')
set_font(r, 18, True, 'C00000')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
r = p.add_run('ТЕКСЕРУ ЖҮРГІЗІЛМЕДІ  ҚҰЖАТ ТЕК ФОРМАТТЫ КӨРСЕТЕДІ')
set_font(r, 11, True, 'C00000')

p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Мәтіндік сәйкестікті тексеру туралы анықтама үлгісі')
set_font(r, 16, True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(16)
r = p.add_run('Ғылыми жоба құжатының демонстрациялық нысаны')
set_font(r, 12)

meta = [
    ('Жоба атауы', 'Бабалар үні цифрлық әлемде'),
    ('Толық атауы', 'Қазақтың ұлттық музыкалық аспаптарының интерактивті онлайн энциклопедиясы'),
    ('Тексеру жүйесі', 'DEMO CHECK оқу жүйесі  нақты сервис емес'),
    ('Тексеру күні', '____  ____________ 2026 жыл'),
    ('Файл атауы', '____________________________________________'),
]

table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Table Grid'
table.autofit = False
table.columns[0].width = Cm(5.2)
table.columns[1].width = Cm(11.3)
tr_pr = table.rows[0]._tr.get_or_add_trPr()
tbl_header = OxmlElement('w:tblHeader')
tbl_header.set(qn('w:val'), 'true')
tr_pr.append(tbl_header)
for i, text in enumerate(['Көрсеткіш', 'Үлгілік мәлімет']):
    cell = table.rows[0].cells[i]
    shade(cell, 'D9EAF7')
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, 11, True)

for label, value in meta:
    cells = table.add_row().cells
    for c in cells:
        set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    r = cells[0].paragraphs[0].add_run(label)
    set_font(r, 11, True)
    r = cells[1].paragraphs[0].add_run(value)
    set_font(r, 11)

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
r = p.add_run('Үлгілік түпнұсқалық көрсеткіші')
set_font(r, 12, True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
r = p.add_run('100%')
set_font(r, 30, True, '1F6B45')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
r = p.add_run('Бұл сан нақты тексерістен алынған жоқ және тек анықтама үлгісінің сыртқы көрінісін көрсету үшін жазылды.')
set_font(r, 11, True, 'C00000')

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run('Үлгілік қорытынды. ')
set_font(r, 12, True)
r = p.add_run('Ғылыми жоба мәтінінің түпнұсқалық көрсеткіші арнайы антиплагиат жүйесінің нақты есебі алынғаннан кейін ғана осы жолға енгізіледі. Ресми анықтамаға жүйенің атауы, тексеру күні, нақты пайыз, есеп нөмірі және тексерушінің қолы жазылуы керек.')
set_font(r, 12)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
r = p.add_run('Нақты тексерістен кейін толтырылатын жолдар')
set_font(r, 12, True)

for text in [
    'Нақты түпнұсқалық көрсеткіші: __________ %',
    'Мәтіндік сәйкестік көрсеткіші: __________ %',
    'Есептің нөмірі немесе сілтемесі: ________________________________',
    'Тексерген тұлға: ______________________________________________',
    'Қолы және күні: ______________________________________________',
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.7)
    r = p.add_run(text)
    set_font(r, 12)

footer = sec.footer
p = footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('ҮЛГІ  РЕСМИ ЕМЕС  НАҚТЫ АНТИПЛАГИАТ ЕСЕБІН АЛМАСТЫРМАЙДЫ')
set_font(r, 9, True, 'C00000')

doc.core_properties.title = 'Мәтіндік сәйкестікті тексеру туралы анықтама үлгісі'
doc.core_properties.subject = 'Ресми емес демонстрациялық шаблон'
doc.core_properties.author = 'Template'
doc.save(OUT)
print(OUT)
